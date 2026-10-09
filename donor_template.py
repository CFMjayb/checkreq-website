"""
donor_template.py -- Beacon Donor Management: the Excel upload template for people and households (MG-01).

A parish that switches to Beacon loads its people from one standard workbook:
  template_build()                      the blank workbook (instructions, People, Households, Lists, one example row)
  template_validate(ctx, file_bytes)    check EVERY row and report errors by sheet, row and column. Writes nothing.
  template_import(ctx, file_bytes, dry_run=True, skip_duplicates=True)
                                        dry run by default (the same report plus what would be created). A real
                                        import runs all-or-nothing in ONE transaction and only when the file has
                                        no errors: "checked row by row before anything loads".

The membership columns on the People sheet (status, how joined, join date) belong to the membership sheet
of the full template (MG-01 lists gifts and pledges as further sheets, built with Phase 2 and 3).

Rules the checker applies, by column: names and dates, e-mail and phone formats, allowed values, household
keys that exist, one primary contact per household and never a child, spouse keys that exist and pair
people only once, joint statements that have a spouse, member status codes that exist at THIS parish,
and likely duplicates (against the file itself and against people already at this parish). Dates must be
real Excel dates or MM/DD/YYYY text. Example rows (a Row key starting with EXAMPLE) are ignored.
"""
from __future__ import annotations

import datetime as dt
import io
from collections import Counter, defaultdict

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

import db
from donor_core import (
    CONNECTION_KINDS, Ctx, GENDERS, GRADES, HOUSEHOLD_POSITIONS, HOW_JOINED, InvalidInput, MARITAL_STATUSES,
    PermissionDenied, STATEMENT_OPTIONS, check_enum, clean_email, clean_phone, clean_text, log_change, need_people,
    parse_date, to_bool, tx, today,
)
import donor_households as H
import donor_membership as M
import donor_people as P

TEMPLATE_VERSION = "Beacon Donor Management upload template v2"
MAX_BYTES = 8 * 1024 * 1024
MAX_PEOPLE = 5000

# (key, header, width, kind, allowed) -- kind: text | date | bool | enum | email | phone
PEOPLE_COLUMNS = [
    ("row_key", "Row key", 12, "text", None),
    ("record_type", "Record type", 13, "enum", ("Person", "Organization")),
    ("title", "Title", 8, "text", None),
    ("first_name", "First name", 16, "text", None),
    ("middle_name", "Middle name", 14, "text", None),
    ("last_name", "Last name", 18, "text", None),
    ("suffix", "Suffix", 8, "text", None),
    ("goes_by", "Goes by", 14, "text", None),
    ("former_name", "Former name", 16, "text", None),
    ("org_name", "Organization name", 24, "text", None),
    ("gender", "Gender", 11, "enum", ("Female", "Male", "Nonbinary", "Unknown")),
    ("marital_status", "Marital status", 14, "enum", ("Single", "Married", "Widowed", "Divorced", "Separated", "Unknown")),
    ("birth_date", "Birth date", 13, "date", None),
    ("wedding_date", "Wedding date", 13, "date", None),
    ("deceased_date", "Deceased date", 13, "date", None),
    ("email", "Email (preferred)", 28, "email", None),
    ("email2", "Other email", 28, "email", None),
    ("phone_cell", "Cell phone", 15, "phone", None),
    ("phone_home", "Home phone", 15, "phone", None),
    ("phone_work", "Work phone", 15, "phone", None),
    ("household_key", "Household key", 14, "text", None),
    ("household_position", "Household position", 18, "enum", ("Primary Adult", "Secondary Adult", "Child")),
    ("is_primary_contact", "Primary contact", 14, "bool", None),
    ("spouse_key", "Spouse row key", 14, "text", None),
    ("connection", "Connection", 12, "enum", ("Member", "Giver", "Visitor")),
    ("member_status", "Member status code", 18, "text", None),
    ("how_joined", "How joined", 14, "enum", ("Baptism", "Transfer", "Confirmation", "Reception", "Reaffirmation", "Other")),
    ("join_date", "Join date", 13, "date", None),
    ("envelope_number", "Envelope number", 14, "text", None),
    ("statement", "Statement", 12, "enum", ("Individual", "Joint", "None")),
    ("in_directory", "In directory", 12, "bool", None),
    ("do_not_mail", "Do not mail", 12, "bool", None),
    ("do_not_call", "Do not call", 12, "bool", None),
    ("do_not_email", "Do not email", 12, "bool", None),
    # Added in template v2. A v1 workbook without these columns still loads (see OPTIONAL_PEOPLE_KEYS).
    ("alt_name", "Alt name", 14, "text", None),
    ("occupation", "Occupation", 18, "text", None),
    ("employer", "Employer", 20, "text", None),
    ("school", "School", 20, "text", None),
    ("grade", "Grade", 10, "enum", ("Pre-K", "K") + tuple(str(n) for n in range(1, 13)) + ("College", "Graduate")),
]
OPTIONAL_PEOPLE_KEYS = frozenset({"alt_name", "occupation", "employer", "school", "grade"})
HOUSEHOLD_COLUMNS = [
    ("household_key", "Household key", 14, "text", None),
    ("name", "Household name", 26, "text", None),
    ("salutation", "Salutation", 22, "text", None),
    ("directory_name", "Directory name", 24, "text", None),
    ("address1", "Address", 28, "text", None),
    ("address2", "Address line 2", 18, "text", None),
    ("city", "City", 16, "text", None),
    ("state", "State", 8, "text", None),
    ("postal_code", "Zip", 10, "text", None),
    ("home_phone", "Home phone", 15, "phone", None),
    ("mail_address1", "Mailing address (if different)", 28, "text", None),
    ("mail_address2", "Mailing line 2", 18, "text", None),
    ("mail_city", "Mailing city", 16, "text", None),
    ("mail_state", "Mailing state", 8, "text", None),
    ("mail_postal_code", "Mailing zip", 10, "text", None),
]

_NAVY, _BLUE = "1F4E79", "2E75B6"


def _norm(s) -> str:
    return " ".join(str(s or "").strip().lower().split())


# ── Build ───────────────────────────────────────────────────────────────────────────────────────
def template_build() -> bytes:
    wb = Workbook()
    head_font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
    head_fill = PatternFill("solid", fgColor=_NAVY)
    note_fill = PatternFill("solid", fgColor="FFF4CC")
    thin = Side(style="thin", color="BFBFBF")
    body = Font(name="Arial", size=10)

    ws = wb.active
    ws.title = "Instructions"
    lines = [
        (TEMPLATE_VERSION, True),
        ("", False),
        ("How to use this workbook", True),
        ("1. Fill in the People sheet (one row per person or organization) and the Households sheet (one row per household).", False),
        ("2. Give every person a Row key you make up (P1, P2, ...). Give every household a Household key (H1, H2, ...). Put the same Household key on the people who live together.", False),
        ("3. Link spouses by putting each spouse's Row key in the other's 'Spouse row key' column (one side is enough).", False),
        ("4. Dates: type them as real dates (MM/DD/YYYY). Yes/No columns: Yes or No (blank means No, except 'In directory', where blank means Yes).", False),
        ("5. Member status code: use one of this parish's codes (see the Lists sheet, or Parish Settings in Beacon). Leave blank for a giver or visitor.", False),
        ("6. Upload the file in Beacon under People > Import. Beacon checks every row first and shows each problem by sheet, row and column. Nothing is loaded until the whole file is clean.", False),
        ("7. The first upload is a DRY RUN: it only reports what would be created. Run it again with 'Load for real' when the report is clean.", False),
        ("", False),
        ("Rules worth knowing", True),
        ("- Do not rename the sheets or change the column headings in row 1.", False),
        ("- A household can have only one Primary contact, and a child cannot be the primary contact.", False),
        ("- A person can have only one spouse. 'Joint' statements need a spouse.", False),
        ("- A person who is already at this parish with the same name and birth date is reported as a possible duplicate and skipped (you can choose to create them anyway).", False),
        ("- People under 18 are loaded like anyone else, and are left out of directories and exports by default.", False),
        ("- The example row (Row key EXAMPLE1) is ignored. Delete it or leave it.", False),
        ("- Use made-up data only when testing. Real people belong in the file only when you are loading a real parish.", False),
    ]
    for i, (text, bold) in enumerate(lines, start=1):
        c = ws.cell(row=i, column=1, value=text)
        c.font = Font(name="Arial", size=11 if bold else 10, bold=bold, color=_NAVY if bold else "000000")
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 120

    def sheet(title, cols, example):
        s = wb.create_sheet(title)
        for j, (_, header, width, kind, allowed) in enumerate(cols, start=1):
            cell = s.cell(row=1, column=j, value=header)
            cell.font, cell.fill = head_font, head_fill
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            s.column_dimensions[get_column_letter(j)].width = width
        s.row_dimensions[1].height = 32
        s.freeze_panes = "A2"
        for j, ((key, _, _, kind, allowed), val) in enumerate(zip(cols, example), start=1):
            cell = s.cell(row=2, column=j, value=val)
            cell.fill = note_fill
            cell.font = Font(name="Arial", size=10, italic=True)
            if kind == "date":
                cell.number_format = "mm/dd/yyyy"
        for j, (key, _, _, kind, allowed) in enumerate(cols, start=1):
            letter = get_column_letter(j)
            rng = f"{letter}2:{letter}{MAX_PEOPLE + 1}"
            for r in range(3, 40):
                if kind == "date":
                    s.cell(row=r, column=j).number_format = "mm/dd/yyyy"
                s.cell(row=r, column=j).font = body
            if kind == "date":
                for r in range(40, 400):
                    s.cell(row=r, column=j).number_format = "mm/dd/yyyy"
            if kind == "enum" and allowed:
                dv = DataValidation(type="list", formula1='"' + ",".join(allowed) + '"', allow_blank=True)
                dv.error, dv.errorTitle, dv.showErrorMessage = "Pick a value from the list.", "Not allowed", True
                s.add_data_validation(dv)
                dv.add(rng)
            if kind == "bool":
                dv = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
                dv.error, dv.errorTitle, dv.showErrorMessage = "Yes or No.", "Not allowed", True
                s.add_data_validation(dv)
                dv.add(rng)
        return s

    sheet("People", PEOPLE_COLUMNS, [
        "EXAMPLE1", "Person", "Mrs.", "Pat", "Q.", "Examplesmith", "", "Patty", "", "", "Female", "Married",
        dt.date(1960, 4, 12), dt.date(1985, 6, 18), None, "pat.examplesmith@example.org", "", "(410) 555-0100",
        "", "", "H1", "Primary Adult", "Yes", "", "Member", "MEMBER", "Transfer", dt.date(2010, 3, 2), "101",
        "Individual", "Yes", "No", "No", "No", "", "Retired teacher", "", "", ""])
    sheet("Households", HOUSEHOLD_COLUMNS, [
        "H1", "The Examplesmith Household", "Pat & Sam", "Examplesmith, Pat & Sam", "100 Example Lane", "", "Exampleville",
        "MD", "21000", "(410) 555-0101", "", "", "", "", ""])
    lists = wb.create_sheet("Lists")
    lists.cell(row=1, column=1, value="Allowed values (the drop-downs on the other sheets come from here)").font = Font(name="Arial", bold=True)
    col = 1
    for title, vals in (("Record type", ("Person", "Organization")), ("Gender", ("Female", "Male", "Nonbinary", "Unknown")),
                        ("Marital status", ("Single", "Married", "Widowed", "Divorced", "Separated", "Unknown")),
                        ("Household position", ("Primary Adult", "Secondary Adult", "Child")),
                        ("Grade", ("Pre-K", "K") + tuple(str(n) for n in range(1, 13)) + ("College", "Graduate")),
                        ("Connection", ("Member", "Giver", "Visitor")),
                        ("How joined", ("Baptism", "Transfer", "Confirmation", "Reception", "Reaffirmation", "Other")),
                        ("Statement", ("Individual", "Joint", "None"))):
        lists.cell(row=3, column=col, value=title).font = Font(name="Arial", bold=True, color=_BLUE)
        for i, v in enumerate(vals, start=4):
            lists.cell(row=i, column=col, value=v)
        lists.column_dimensions[get_column_letter(col)].width = 20
        col += 1
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def template_filename() -> str:
    return f"Beacon People Upload Template {today().strftime('%Y-%m-%d')}.xlsx"


# ── Read and check ──────────────────────────────────────────────────────────────────────────────
class _Report:
    def __init__(self):
        self.errors: list[dict] = []
        self.warnings: list[dict] = []

    def error(self, sheet, row, column, message):
        self.errors.append({"sheet": sheet, "row": row, "column": column, "message": message})

    def warn(self, sheet, row, column, message):
        self.warnings.append({"sheet": sheet, "row": row, "column": column, "message": message})


def _read_sheet(wb, name, cols, rep: _Report, optional=frozenset()):
    """Rows as {key: cell value, '_row': excel row}, matching headers case-insensitively. Missing or extra
    sheets and headers are reported, never guessed. A column named in `optional` may be absent (a workbook built
    from an earlier template version): its values are simply blank."""
    if name not in wb.sheetnames:
        rep.error(name, None, None, f"The sheet '{name}' is missing. Start from the Beacon template.")
        return []
    ws = wb[name]
    header = {_norm(c.value): idx for idx, c in enumerate(next(ws.iter_rows(min_row=1, max_row=1)), start=0) if c.value}
    index = {}
    for key, head, *_ in cols:
        if _norm(head) not in header:
            if key not in optional:
                rep.error(name, 1, head, f"The column '{head}' is missing from row 1. Start from the Beacon template.")
        else:
            index[key] = header[_norm(head)]
    rows = []
    for rnum, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if row is None or all(v is None or (isinstance(v, str) and not v.strip()) for v in row):
            continue
        rec = {k: (row[i] if i < len(row) else None) for k, i in index.items()}
        rec["_row"] = rnum
        rows.append(rec)
    return rows


def _val(rec, key):
    v = rec.get(key)
    if isinstance(v, str):
        v = v.strip()
        return v or None
    return v


def _date_cell(v, field):
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        raise InvalidInput(f"{field} must be a real date (MM/DD/YYYY), not a number.", field)
    return parse_date(v, field=field.lower(), allow_future=(field == "Wedding date"))


def _load_workbook(file_bytes: bytes, rep: _Report):
    if not file_bytes:
        rep.error("File", None, None, "The file is empty.")
        return None
    if len(file_bytes) > MAX_BYTES:
        rep.error("File", None, None, f"The file is larger than {MAX_BYTES // (1024 * 1024)} MB.")
        return None
    try:
        return load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=False)
    except Exception:
        rep.error("File", None, None, "That is not an Excel workbook (.xlsx). Start from the Beacon template.")
        return None


def _check(ctx: Ctx, file_bytes: bytes) -> dict:
    """Parse and validate. Returns {report, people, households, cleaned...} with no writes."""
    need_people(ctx, "people.edit", "Only clergy, membership editors and parish admins can load people.")
    rep = _Report()
    wb = _load_workbook(file_bytes, rep)
    result = {"rep": rep, "people": [], "households": {}, "skipped_example": 0}
    if wb is None:
        return result
    people_rows = _read_sheet(wb, "People", PEOPLE_COLUMNS, rep, optional=OPTIONAL_PEOPLE_KEYS)
    hh_rows = _read_sheet(wb, "Households", HOUSEHOLD_COLUMNS, rep)
    if len(people_rows) > MAX_PEOPLE:
        rep.error("People", None, None, f"More than {MAX_PEOPLE} rows. Split the file.")
        return result

    codes = {}
    if ctx.can("membership.view"):
        for c in M.status_code_list(ctx):
            codes[c["code"].upper()] = c
    # households
    households: dict[str, dict] = {}
    for rec in hh_rows:
        r = rec["_row"]
        key = _val(rec, "household_key")
        if key is None:
            rep.error("Households", r, "Household key", "Every household needs a Household key.")
            continue
        if str(key).upper().startswith("EXAMPLE"):
            result["skipped_example"] += 1
            continue
        if key in households:
            rep.error("Households", r, "Household key", f"The key '{key}' is used twice (first on row {households[key]['_row']}).")
            continue
        data = {}
        for k, head, w, kind, allowed in HOUSEHOLD_COLUMNS[1:]:
            v = _val(rec, k)
            try:
                data[k] = clean_phone(v, field=head) if kind == "phone" else clean_text(v, field=head, max_len=160)
            except InvalidInput as e:
                rep.error("Households", r, head, e.message)
        households[str(key)] = {"_row": r, "data": data, "members": []}
    result["households"] = households

    # people
    seen_keys: dict[str, int] = {}
    cleaned: list[dict] = []
    for rec in people_rows:
        r = rec["_row"]
        key = _val(rec, "row_key")
        if key is None:
            rep.error("People", r, "Row key", "Every row needs a Row key (P1, P2, ...).")
            continue
        key = str(key)
        if key.upper().startswith("EXAMPLE"):
            result["skipped_example"] += 1
            continue
        if key in seen_keys:
            rep.error("People", r, "Row key", f"The key '{key}' is used twice (first on row {seen_keys[key]}).")
            continue
        seen_keys[key] = r
        p = {"_row": r, "key": key, "fields": {}, "contacts": [], "errors": 0}

        def bad(col_head, msg):
            rep.error("People", r, col_head, msg)
            p["errors"] += 1

        rt = _val(rec, "record_type")
        try:
            record_type = check_enum(rt or "Person", ("person", "organization"), field="Record type", allow_blank=False)
        except InvalidInput as e:
            bad("Record type", e.message)
            record_type = "person"
        p["record_type"] = record_type
        for k, head, w, kind, allowed in PEOPLE_COLUMNS:
            if k in ("row_key", "record_type") or k in ("email", "email2", "phone_cell", "phone_home", "phone_work"):
                continue
            if k in ("household_key", "household_position", "is_primary_contact", "spouse_key", "connection",
                     "member_status", "how_joined", "join_date", "envelope_number", "statement"):
                continue
            v = _val(rec, k)
            try:
                if kind == "date":
                    d = _date_cell(rec.get(k), head)
                    if d is not None:
                        p["fields"][k] = d
                elif kind == "enum":
                    if isinstance(v, float) and v.is_integer():
                        v = int(v)                      # Excel stores a typed 5 as 5.0: a grade of 5, not "5.0"
                    if v is not None:
                        p["fields"][k] = check_enum(v, {"gender": GENDERS, "marital_status": MARITAL_STATUSES, "grade": GRADES}[k], field=head)
                elif kind == "bool":
                    if v is not None:
                        p["fields"][k] = to_bool(v, field=head)
                else:
                    if v is not None:
                        p["fields"][k] = clean_text(v, field=head, max_len=160)
            except InvalidInput as e:
                bad(head, e.message)
        try:
            P.clean_person_fields(p["fields"], record_type, creating=True)
        except InvalidInput as e:
            bad({"last_name": "Last name", "org_name": "Organization name"}.get(e.field, "Last name"), e.message)
        for k, head, preferred in (("email", "Email (preferred)", True), ("email2", "Other email", False)):
            v = _val(rec, k)
            if v is None:
                continue
            try:
                p["contacts"].append({"kind": "email", "value": clean_email(v, field=head), "is_preferred": preferred, "subtype": "other"})
            except InvalidInput as e:
                bad(head, e.message)
        for k, head, sub, preferred in (("phone_cell", "Cell phone", "cell", True), ("phone_home", "Home phone", "home", False),
                                        ("phone_work", "Work phone", "work", False)):
            v = _val(rec, k)
            if v is None:
                continue
            try:
                p["contacts"].append({"kind": "phone", "value": clean_phone(str(v), field=head), "subtype": sub, "is_preferred": False})
            except InvalidInput as e:
                bad(head, e.message)
        # household
        hk = _val(rec, "household_key")
        p["household_key"] = str(hk) if hk is not None else None
        pos = _val(rec, "household_position")
        p["position"] = None
        try:
            if pos is not None:
                p["position"] = check_enum(pos, HOUSEHOLD_POSITIONS, field="Household position")
        except InvalidInput as e:
            bad("Household position", e.message)
        try:
            p["primary"] = to_bool(_val(rec, "is_primary_contact"), field="Primary contact")
        except InvalidInput as e:
            bad("Primary contact", e.message)
            p["primary"] = False
        if p["household_key"]:
            if p["household_key"] not in households:
                bad("Household key", f"There is no household '{p['household_key']}' on the Households sheet.")
            else:
                if p["position"] is None:
                    p["position"] = "primary_adult"
                households[p["household_key"]]["members"].append(p)
        elif p["position"] or p["primary"]:
            bad("Household key", "A household position or primary contact needs a Household key.")
        if record_type == "organization" and p["household_key"]:
            bad("Household key", "An organization does not belong to a household.")
        p["spouse_key"] = _val(rec, "spouse_key")
        p["spouse_key"] = str(p["spouse_key"]) if p["spouse_key"] is not None else None
        # connection and membership
        try:
            p["connection"] = check_enum(_val(rec, "connection") or ("member" if ctx.can("membership.edit") else "giver"),
                                         CONNECTION_KINDS, field="Connection", allow_blank=False)
        except InvalidInput as e:
            bad("Connection", e.message)
            p["connection"] = "giver"
        ms = _val(rec, "member_status")
        p["member_status"] = str(ms).upper() if ms is not None else None
        if p["member_status"]:
            if not ctx.can("membership.edit"):
                bad("Member status code", "You do not have the membership role, so member status cannot be loaded.")
            elif p["member_status"] not in codes:
                bad("Member status code", f"'{ms}' is not one of this parish's codes. Codes: {', '.join(sorted(codes)) or 'none yet'}.")
        if p["connection"] == "member" and not ctx.can("membership.edit"):
            bad("Connection", "Only clergy and membership editors can load members. Use Giver or Visitor.")
        hj = _val(rec, "how_joined")
        p["how_joined"] = None
        try:
            if hj is not None:
                p["how_joined"] = check_enum(hj, HOW_JOINED, field="How joined")
        except InvalidInput as e:
            bad("How joined", e.message)
        try:
            p["join_date"] = _date_cell(rec.get("join_date"), "Join date")
        except InvalidInput as e:
            bad("Join date", e.message)
            p["join_date"] = None
        if (p["how_joined"] or p["join_date"]) and not p["member_status"]:
            bad("Member status code", "How joined and join date need a member status code.")
        env = _val(rec, "envelope_number")
        p["envelope"] = str(env) if env is not None else None
        try:
            if p["envelope"]:
                clean_text(p["envelope"], field="Envelope number", max_len=20)
        except InvalidInput as e:
            bad("Envelope number", e.message)
        try:
            p["statement"] = check_enum(_val(rec, "statement"), STATEMENT_OPTIONS, field="Statement")
        except InvalidInput as e:
            bad("Statement", e.message)
            p["statement"] = None
        cleaned.append(p)         # (the privacy columns were read with the other person fields above)
    by_key = {p["key"]: p for p in cleaned}

    # households: one primary contact, never a child, and warn on empty ones
    for hk, h in households.items():
        prim = [m for m in h["members"] if m.get("primary")]
        if len(prim) > 1:
            for m in prim[1:]:
                rep.error("People", m["_row"], "Primary contact", f"Household '{hk}' already has a primary contact (row {prim[0]['_row']}).")
        for m in prim:
            if m.get("position") == "child":
                rep.error("People", m["_row"], "Primary contact", "A child cannot be the primary contact.")
        if not h["members"]:
            rep.warn("Households", h["_row"], "Household key", f"No one on the People sheet belongs to household '{hk}'. It will be created empty.")
    # spouses
    spouse_of: dict[str, str] = {}
    pair_rows: dict[str, int] = {}
    for p in cleaned:
        sk = p["spouse_key"]
        if not sk:
            continue
        if sk not in by_key:
            rep.error("People", p["_row"], "Spouse row key", f"There is no person with the Row key '{sk}'.")
            continue
        if sk == p["key"]:
            rep.error("People", p["_row"], "Spouse row key", "A person cannot be their own spouse.")
            continue
        if by_key[sk]["record_type"] != "person" or p["record_type"] != "person":
            rep.error("People", p["_row"], "Spouse row key", "An organization cannot have a spouse.")
            continue
        for a, b in ((p["key"], sk), (sk, p["key"])):
            if a in spouse_of and spouse_of[a] != b:
                rep.error("People", p["_row"], "Spouse row key", f"'{a}' would have two spouses ('{spouse_of[a]}' and '{b}').")
            spouse_of.setdefault(a, b)
    for p in cleaned:
        if p["statement"] == "joint" and p["key"] not in spouse_of:
            rep.error("People", p["_row"], "Statement", "A joint statement needs a Spouse row key (on this row or on the spouse's row).")
    # duplicates: within the file, and against people already at this parish
    sig = defaultdict(list)
    for p in cleaned:
        f = p["fields"]
        if p["record_type"] == "person" and f.get("first_name") and f.get("last_name"):
            sig[(_norm(f["first_name"]), _norm(f["last_name"]), f.get("birth_date"))].append(p)
        elif p["record_type"] == "organization" and f.get("org_name"):
            sig[("org", _norm(f["org_name"]), None)].append(p)
    for s, group in sig.items():
        if len(group) > 1:
            for dup in group[1:]:
                rep.warn("People", dup["_row"], "Last name", f"Looks like a duplicate of row {group[0]['_row']} in this file.")
    existing = db.query(
        "SELECT p.id, p.record_type, LOWER(p.first_name) AS fn, LOWER(p.last_name) AS ln, LOWER(p.org_name) AS org, p.birth_date "
        "FROM donor.person p JOIN donor.parish_connection pc ON pc.person_id = p.id AND pc.parish_id = %s", (ctx.parish_id,))
    ex_person = {(r["fn"], r["ln"], r["birth_date"]): r["id"] for r in existing if r["record_type"] == "person"}
    ex_person_anydob = defaultdict(list)
    for r in existing:
        if r["record_type"] == "person":
            ex_person_anydob[(r["fn"], r["ln"])].append(r)
    ex_org = {r["org"]: r["id"] for r in existing if r["record_type"] == "organization"}
    for p in cleaned:
        f = p["fields"]
        match = None
        if p["record_type"] == "organization":
            match = ex_org.get(_norm(f.get("org_name")))
        elif f.get("first_name") and f.get("last_name"):
            k = (_norm(f["first_name"]), _norm(f["last_name"]))
            if (k[0], k[1], f.get("birth_date")) in ex_person:
                match = ex_person[(k[0], k[1], f.get("birth_date"))]
            else:
                for r in ex_person_anydob.get(k, []):
                    if r["birth_date"] is None or f.get("birth_date") is None:
                        match = r["id"]
                        break
        p["duplicate_of"] = match
        if match:
            rep.warn("People", p["_row"], "Last name", f"Already at this parish (person #{match}). It will be skipped unless you choose to create duplicates.")
    # envelope numbers shared by different households
    env_hh = defaultdict(set)
    for p in cleaned:
        if p["envelope"]:
            env_hh[p["envelope"]].add(p["household_key"] or f"row{p['_row']}")
    for env, hhs in env_hh.items():
        if len(hhs) > 1:
            first = next(p for p in cleaned if p["envelope"] == env)
            rep.warn("People", first["_row"], "Envelope number", f"Envelope {env} is used by more than one household. Families can share one, strangers should not.")
    result["people"] = cleaned
    return result


def _households_to_create(res: dict, kept_people: list[dict]) -> list[str]:
    """Household keys to create: those with at least one person being loaded, plus any the file lists with no one
    in them at all. A household whose every member is a skipped duplicate is not created again."""
    used = {p["household_key"] for p in kept_people if p.get("household_key")}
    return [hk for hk, h in res["households"].items() if hk in used or not h["members"]]


def _summarize(res: dict, skip_duplicates: bool) -> dict:
    rep: _Report = res["rep"]
    people = res["people"]
    dup = [p for p in people if p.get("duplicate_of")]
    would_create = [p for p in people if not (skip_duplicates and p.get("duplicate_of"))]
    return {
        "ok": not rep.errors,
        "errors": rep.errors, "warnings": rep.warnings,
        "people_rows": len(people), "household_rows": len(res["households"]),
        "possible_duplicates": len(dup),
        "would_create_people": len(would_create) if not rep.errors else 0,
        "would_skip_duplicates": len(dup) if skip_duplicates else 0,
        "would_create_households": len(_households_to_create(res, would_create)) if not rep.errors else 0,
        "example_rows_ignored": res["skipped_example"],
    }


def template_validate(ctx: Ctx, file_bytes: bytes, skip_duplicates: bool = True) -> dict:
    """Check every row. Writes nothing. Returns {ok, errors:[{sheet,row,column,message}], warnings, counts}."""
    return _summarize(_check(ctx, file_bytes), skip_duplicates)


def template_import(ctx: Ctx, file_bytes: bytes, dry_run: bool = True, skip_duplicates: bool = True, *, cur=None) -> dict:
    """Dry run by default. A real import happens only when the file has NO errors, and then everything loads in ONE
    transaction: any failure part-way rolls the whole file back, so a parish never ends up half loaded."""
    res = _check(ctx, file_bytes)
    report = _summarize(res, skip_duplicates)
    report["imported"] = False
    report["dry_run"] = dry_run
    if report["errors"] or dry_run:
        return report
    people = [p for p in res["people"] if not (skip_duplicates and p.get("duplicate_of"))]
    created: dict[str, int] = {}
    with tx(cur) as c:
        hh_ids: dict[str, int] = {}
        for hk in _households_to_create(res, people):
            hh_ids[hk] = H.household_create(ctx, res["households"][hk]["data"], None, cur=c)["id"]
        for p in people:
            row = P.person_create(ctx, p["fields"], record_type=p["record_type"], connection_kind=p["connection"],
                                  contacts=_contacts(p), allow_duplicate=True, cur=c)
            pid = row["id"]
            created[p["key"]] = pid
            conn_changes = {}
            if p.get("envelope"):
                conn_changes["envelope_number"] = p["envelope"]
            if conn_changes:
                H.parish_connection_set(ctx, pid, conn_changes, cur=c)
            if p.get("member_status"):
                data = {"status_code": p["member_status"]}
                if p.get("how_joined"):
                    data["how_joined"] = p["how_joined"]
                if p.get("join_date"):
                    data["join_date"] = p["join_date"]
                M.membership_update(ctx, pid, data, cur=c)
            hk = p.get("household_key")
            if hk and hk in hh_ids:
                H.household_add_member(ctx, hh_ids[hk], pid, p.get("position") or "primary_adult", bool(p.get("primary")), cur=c)
        done_pairs = set()
        for p in people:
            sk = p.get("spouse_key")
            if sk and sk in created and p["key"] in created:
                pair = tuple(sorted((p["key"], sk)))
                if pair not in done_pairs:
                    done_pairs.add(pair)
                    H.spouse_link_set(ctx, created[p["key"]], created[sk], None, cur=c)
        for p in people:
            if p.get("statement") and p["key"] in created:
                H.parish_connection_set(ctx, created[p["key"]], {"statement_option": p["statement"]}, cur=c)
        log_change(c, ctx, "import", None, "people", None, f"{len(created)} people, {len(hh_ids)} households",
                   kind="create", scope="parish")
    report.update({"imported": True, "created_people": len(created), "created_households": len(hh_ids), "person_ids": created})
    return report


def _contacts(p: dict) -> list[dict]:
    """The contacts for person_create: the first phone is preferred, e-mail keeps its own flag."""
    out, phones = [], 0
    for c in p["contacts"]:
        c = dict(c)
        if c["kind"] == "phone":
            c["is_preferred"] = (phones == 0)
            phones += 1
        out.append(c)
    return out
